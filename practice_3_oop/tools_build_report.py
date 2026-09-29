from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
REPORTS_DIR = ROOT / "reports"
OUTPUT = REPORTS_DIR / "Отчёт_ПР3_Халин.docx"


def set_font(run, size: int = 12, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def write_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_font(run, 11, bold)
    set_cell_borders(cell)


def setup_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def centered_paragraph(document: Document, text: str, size: int = 12, bold: bool = False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_font(run, size, bold)
    return paragraph


def heading(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"{number}. {text}")
    set_font(run, 13, True)


def body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_font(run, 12)


def bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_font(run, 12)


def data_table(
    document: Document,
    rows: list[tuple[str, str]],
    header: tuple[str, str] | None = None,
) -> None:
    offset = 1 if header else 0
    table = document.add_table(rows=len(rows) + offset, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    if header:
        for index, value in enumerate(header):
            cell = table.cell(0, index)
            write_cell(cell, value, True)
            shade_cell(cell, "D9EAF7")

    for row_index, (left, right) in enumerate(rows, start=offset):
        write_cell(table.cell(row_index, 0), left, True)
        write_cell(table.cell(row_index, 1), right)

    document.add_paragraph()


def title_page(document: Document) -> None:
    for line in [
        "МИНОБРНАУКИ РОССИИ",
        "Федеральное государственное бюджетное образовательное учреждение высшего образования",
        "«МИРЭА - Российский технологический университет»",
        "РТУ МИРЭА",
    ]:
        centered_paragraph(document, line)

    document.add_paragraph()
    centered_paragraph(document, "Практическое занятие № 3")
    centered_paragraph(
        document,
        "по дисциплине «Технологии разработки приложений на базе фреймворков»",
    )
    centered_paragraph(
        document,
        "Тема работы: Объектно-ориентированное программирование в Python. Классы, объекты и взаимодействие объектов",
    )

    document.add_paragraph()
    table = document.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        (
            "Выполнил:",
            "студент группы\nЭФБО-11-24\n(группа)",
            "\n\n(подпись)",
            "\nХалин П.С.\n(Фамилия И.О.)",
        ),
        (
            "Принял:",
            "доц.\n(должность)",
            "\n\n(подпись)",
            "\nЯщун Т.В.\n(Фамилия И.О.)",
        ),
    ]
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            write_cell(table.cell(row_index, col_index), value, col_index == 0)

    for _ in range(8):
        document.add_paragraph()
    centered_paragraph(document, "Москва")
    centered_paragraph(document, "2026")
    document.add_section(WD_SECTION.NEW_PAGE)


def build_report() -> None:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    setup_document(document)
    title_page(document)

    centered_paragraph(document, "ОТЧЁТ", 16, True)
    centered_paragraph(document, "о выполнении практической работы № 3", 14)

    heading(document, 1, "Общие сведения")
    data_table(
        document,
        [
            ("Дисциплина", "Технологии разработки приложений на базе фреймворков"),
            (
                "Номер и тема работы",
                "ПР3. Объектно-ориентированное программирование в Python. Классы, объекты и взаимодействие объектов",
            ),
            ("Группа", "ЭФБО-11-24"),
            ("ФИО студента", "Халин Пётр Сергеевич"),
            ("Год выполнения", "2026"),
        ],
    )

    heading(document, 2, "Выполненные материалы Яндекс Практикума")
    data_table(
        document,
        [
            (
                "Раздел 03. Объекты и классы",
                "Введение в ООП; классы и объекты в Python; собственные классы и объекты; атрибуты класса и объекта; методы объекта; магический метод __str__; практика по теме",
            ),
            (
                "Раздел 03. Знакомство с ООП",
                "Принципы ООП; наследование; полиморфизм; инкапсуляция; практика по теме",
            ),
            (
                "Раздел 03. Расширенные возможности Python",
                "Декораторы; декораторы для методов класса; практика по теме",
            ),
        ],
        ("Раздел", "Выполненные темы"),
    )
    body(
        document,
        "Примечание. Материалы Яндекс Практикума выполняются в личном кабинете платформы под персональной учётной записью и отмечаются студентом в чек-листах методических указаний перед защитой работы.",
    )

    heading(document, 3, "Индивидуальный проект")
    data_table(
        document,
        [
            ("Название проекта", "Сервис поиска мест по интересам"),
            (
                "Основные классы",
                "User (пользователь), Category (категория), Place (место), Recommendation (рекомендация)",
            ),
            (
                "Выполненная доработка",
                "Основные сущности проекта представлены классами; реализованы конструкторы __init__, строковое представление __str__, методы объектов и взаимодействие между объектами.",
            ),
            (
                "Структура модулей",
                "models/ (классы предметной области), services.py (бизнес-логика), storage.py (JSON-хранилище), main.py (демонстрационный запуск), tests/ (автоматизированные тесты)",
            ),
            (
                "Ссылка на Git-репозиторий",
                "https://github.com/pendal4f/Framework-based-application-development",
            ),
        ],
    )

    heading(document, 4, "Результат работы")
    for item in [
        "основные сущности проекта представлены классами и объектами;",
        "для классов User, Category, Place и Recommendation реализованы атрибуты, методы, __init__ и __str__;",
        "организовано взаимодействие объектов: Recommendation хранит объект User и объект Place, а Place связан с объектом Category;",
        "функции ПР2 адаптированы к объектной модели: подбор мест по интересам, фильтрация, сортировка, создание и отмена рекомендаций;",
        "коллекции проекта содержат объекты предметной области;",
        "данные продолжают храниться в JSON, при загрузке преобразуются в объекты, а при сохранении возвращаются в словарный формат;",
        "автоматизированные тесты адаптированы к ООП-модели и проверяют создание объектов, методы, связи объектов и JSON-преобразование;",
        "README.md обновлён с описанием классов, структуры проекта, запуска программы и тестов.",
    ]:
        bullet(document, item)

    heading(document, 5, "Вывод")
    body(
        document,
        "В ходе выполнения ПР3 индивидуальный проект «Сервис поиска мест по интересам» был переработан в объектно-ориентированную модель. Основные данные и операции объединены в классах, взаимодействие между пользователями, категориями, местами и рекомендациями стало более наглядным, а функциональность предыдущих практических работ сохранена. Проект готов к дальнейшему развитию и переходу к веб-приложению на Django.",
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build_report()
